"""
6-sheet Excel workbook generator.
Reuses DB blue styling from oil-reit-analysis.
"""

import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule

# Deutsche Bank styling
DB_BLUE = "0018A8"
DB_LIGHT = "E8EDF2"
DB_MID = "C7D3E3"

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", fgColor=DB_BLUE)
LIGHT_FILL = PatternFill("solid", fgColor=DB_LIGHT)
BOLD = Font(bold=True)
BLUE_FONT = Font(color=DB_BLUE)
BLUE_BOLD = Font(bold=True, color=DB_BLUE)
GREEN_FONT = Font(color="006600")
BLACK_FONT = Font(color="000000")
YELLOW_FILL = PatternFill("solid", fgColor="FFFF00")
RED_FILL_LIGHT = PatternFill("solid", fgColor="FFE0E0")
GREEN_FILL_LIGHT = PatternFill("solid", fgColor="E0FFE0")
THIN_BORDER = Border(bottom=Side(style="thin", color=DB_MID))
PCT_FORMAT = "0.00%"
NUM_FORMAT = "0.0000"
SCORE_FORMAT = "0.0"


def _header_row(ws, row, headers, fill=HEADER_FILL, font=HEADER_FONT):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = font
        c.fill = fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def _auto_width(ws):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[letter].width = min(max_len + 3, 26)


def generate_excel(analysis: dict) -> io.BytesIO:
    """Generate 6-sheet Excel workbook from analysis results."""
    wb = Workbook()

    _build_portfolio_summary(wb, analysis)
    _build_diversification_metrics(wb, analysis)
    _build_correlation_matrix(wb, analysis)
    _build_risk_factor_detail(wb, analysis)
    _build_overlap_analysis(wb, analysis)
    _build_damodaran_reference(wb, analysis)

    # Remove default sheet if it still exists
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _build_portfolio_summary(wb: Workbook, analysis: dict):
    ws = wb.create_sheet("Portfolio Summary", 0)

    risk = analysis.get("risk_factors", {})
    portfolio = analysis.get("portfolio_metrics", {})
    tickers = analysis.get("tickers", [])
    weights = analysis.get("weights", {})
    per_holding = risk.get("per_holding", {})
    weighted = risk.get("portfolio_weighted", {})

    from backend.risk_factors import RISK_FACTOR_LABELS, RISK_FACTORS

    # Title
    ws.cell(row=1, column=1, value="Portfolio Risk Exposure Analysis").font = Font(
        bold=True, size=14, color=DB_BLUE
    )
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)

    # Overall rating
    ws.cell(row=2, column=1, value="Overall Risk Grade:").font = BOLD
    grade_cell = ws.cell(row=2, column=2, value=risk.get("risk_grade", "N/A"))
    grade_cell.font = Font(bold=True, size=14, color="FF0000" if risk.get("risk_grade", "C") in ("D", "F") else DB_BLUE)

    ws.cell(row=2, column=3, value="Overall Score:").font = BOLD
    ws.cell(row=2, column=4, value=risk.get("overall_score", 0)).number_format = SCORE_FORMAT

    # Holdings table
    row = 4
    headers = ["Ticker", "Weight"] + [RISK_FACTOR_LABELS.get(f, f) for f in RISK_FACTORS]
    _header_row(ws, row, headers)

    for i, ticker in enumerate(tickers):
        row = 5 + i
        ws.cell(row=row, column=1, value=ticker).font = BLUE_FONT  # blue = user input
        ws.cell(row=row, column=2, value=weights.get(ticker, 0)).number_format = PCT_FORMAT
        ws.cell(row=row, column=2).font = BLUE_FONT  # blue = user input

        scores = per_holding.get(ticker, {})
        for j, factor in enumerate(RISK_FACTORS):
            cell = ws.cell(row=row, column=3 + j, value=scores.get(factor, 5))
            cell.number_format = "0"
            cell.font = BLACK_FONT
            # Yellow highlight for high/critical scores
            if scores.get(factor, 5) >= 7:
                cell.fill = YELLOW_FILL

        if i % 2 == 0:
            for col in range(1, len(headers) + 1):
                if ws.cell(row=row, column=col).fill == PatternFill():
                    ws.cell(row=row, column=col).fill = LIGHT_FILL

    # Portfolio weighted row
    row = 5 + len(tickers)
    ws.cell(row=row, column=1, value="Portfolio Weighted").font = BOLD
    ws.cell(row=row, column=2, value="=SUM(B5:B{})".format(4 + len(tickers))).number_format = PCT_FORMAT
    ws.cell(row=row, column=2).font = GREEN_FONT  # green = formula

    for j, factor in enumerate(RISK_FACTORS):
        # SUMPRODUCT formula referencing weight and score columns
        weight_range = f"B5:B{4 + len(tickers)}"
        score_col = get_column_letter(3 + j)
        score_range = f"{score_col}5:{score_col}{4 + len(tickers)}"
        cell = ws.cell(
            row=row,
            column=3 + j,
            value=f"=SUMPRODUCT({weight_range},{score_range})",
        )
        cell.number_format = SCORE_FORMAT
        cell.font = GREEN_FONT  # green = formula

    _auto_width(ws)


def _build_diversification_metrics(wb: Workbook, analysis: dict):
    ws = wb.create_sheet("Diversification Metrics")

    portfolio = analysis.get("portfolio_metrics", {})
    tickers = analysis.get("tickers", [])
    weights = analysis.get("weights", {})
    ind_vols = portfolio.get("individual_volatilities", {})
    mrc = portfolio.get("marginal_risk_contributions", {})

    # Title
    ws.cell(row=1, column=1, value="Diversification Metrics").font = Font(
        bold=True, size=14, color=DB_BLUE
    )

    # Key metrics
    metrics = [
        ("HHI (Herfindahl-Hirschman)", portfolio.get("hhi", 0), NUM_FORMAT),
        ("N_eff (Effective # of Bets)", portfolio.get("n_eff", 0), "0.00"),
        ("Diversification Ratio", portfolio.get("diversification_ratio", 0), "0.00"),
        ("Portfolio Volatility (Ann.)", portfolio.get("portfolio_volatility", 0), PCT_FORMAT),
        ("Distance to Diversification", portfolio.get("distance_to_diversification", 0), "0.00"),
    ]

    row = 3
    _header_row(ws, row, ["Metric", "Value"])
    for i, (name, val, fmt) in enumerate(metrics):
        r = row + 1 + i
        ws.cell(row=r, column=1, value=name).font = BOLD
        c = ws.cell(row=r, column=2, value=val)
        c.number_format = fmt
        c.font = BLACK_FONT
        if i % 2 == 0:
            ws.cell(row=r, column=1).fill = LIGHT_FILL
            c.fill = LIGHT_FILL

    # Per-stock breakdown
    row = 3 + len(metrics) + 3
    _header_row(ws, row, ["Ticker", "Weight", "Ann. Volatility", "Marginal Risk Contribution", "MRC %"])

    for i, ticker in enumerate(tickers):
        r = row + 1 + i
        ws.cell(row=r, column=1, value=ticker).font = BLUE_FONT
        ws.cell(row=r, column=2, value=weights.get(ticker, 0)).number_format = PCT_FORMAT
        ws.cell(row=r, column=3, value=ind_vols.get(ticker, 0)).number_format = PCT_FORMAT

        mrc_val = mrc.get(ticker, 0)
        ws.cell(row=r, column=4, value=mrc_val).number_format = PCT_FORMAT

        port_vol = portfolio.get("portfolio_volatility", 1)
        mrc_pct = mrc_val / port_vol if port_vol > 0 else 0
        ws.cell(row=r, column=5, value=mrc_pct).number_format = PCT_FORMAT

        if i % 2 == 0:
            for col in range(1, 6):
                ws.cell(row=r, column=col).fill = LIGHT_FILL

    # Total MRC row
    r = row + 1 + len(tickers)
    ws.cell(row=r, column=1, value="Total").font = BOLD
    total_col = get_column_letter(4)
    ws.cell(
        row=r,
        column=4,
        value=f"=SUM({total_col}{row+1}:{total_col}{row+len(tickers)})",
    ).font = GREEN_FONT

    _auto_width(ws)


def _build_correlation_matrix(wb: Workbook, analysis: dict):
    ws = wb.create_sheet("Correlation Matrix")

    corr = analysis.get("portfolio_metrics", {}).get("correlation_matrix", {})
    tickers = analysis.get("tickers", [])

    ws.cell(row=1, column=1, value="Pairwise Correlation Matrix").font = Font(
        bold=True, size=14, color=DB_BLUE
    )

    row = 3
    # Header row with ticker names
    for j, t in enumerate(tickers):
        ws.cell(row=row, column=j + 2, value=t).font = HEADER_FONT
        ws.cell(row=row, column=j + 2).fill = HEADER_FILL
        ws.cell(row=row, column=j + 2).alignment = Alignment(horizontal="center")

    for i, t1 in enumerate(tickers):
        r = row + 1 + i
        ws.cell(row=r, column=1, value=t1).font = BOLD
        for j, t2 in enumerate(tickers):
            val = corr.get(t1, {}).get(t2, 0)
            cell = ws.cell(row=r, column=j + 2, value=val)
            cell.number_format = "0.00"
            cell.alignment = Alignment(horizontal="center")

    # Conditional formatting: color scale (green=low, white=mid, red=high)
    if tickers:
        n = len(tickers)
        start = f"B{row+1}"
        end = f"{get_column_letter(n+1)}{row+n}"
        ws.conditional_formatting.add(
            f"{start}:{end}",
            ColorScaleRule(
                start_type="num", start_value=-1, start_color="63BE7B",
                mid_type="num", mid_value=0, mid_color="FFFFFF",
                end_type="num", end_value=1, end_color="F8696B",
            ),
        )

    _auto_width(ws)


def _build_risk_factor_detail(wb: Workbook, analysis: dict):
    ws = wb.create_sheet("Risk Factor Detail")

    risk = analysis.get("risk_factors", {})
    tickers = analysis.get("tickers", [])
    per_holding = risk.get("per_holding", {})
    weighted = risk.get("portfolio_weighted", {})
    benchmark = risk.get("benchmark", {})

    from backend.risk_factors import RISK_FACTOR_LABELS, RISK_FACTORS

    ws.cell(row=1, column=1, value="Risk Factor Detail — 10 Factors × N Holdings").font = Font(
        bold=True, size=14, color=DB_BLUE
    )

    # Factor rows, ticker columns
    row = 3
    headers = ["Risk Factor"] + tickers + ["Portfolio Weighted", "S&P 500 Benchmark", "vs. Benchmark"]
    _header_row(ws, row, headers)

    for i, factor in enumerate(RISK_FACTORS):
        r = row + 1 + i
        label = RISK_FACTOR_LABELS.get(factor, factor)
        ws.cell(row=r, column=1, value=label).font = BOLD

        for j, ticker in enumerate(tickers):
            score = per_holding.get(ticker, {}).get(factor, 5)
            cell = ws.cell(row=r, column=2 + j, value=score)
            cell.number_format = "0"
            cell.alignment = Alignment(horizontal="center")
            if score >= 7:
                cell.fill = YELLOW_FILL

        # Portfolio weighted
        pw = weighted.get(factor, 5)
        pw_cell = ws.cell(row=r, column=2 + len(tickers), value=pw)
        pw_cell.number_format = SCORE_FORMAT
        pw_cell.font = GREEN_FONT

        # Benchmark
        bm = benchmark.get(factor, 5)
        ws.cell(row=r, column=3 + len(tickers), value=bm).number_format = "0"

        # vs benchmark
        diff = pw - bm
        diff_cell = ws.cell(row=r, column=4 + len(tickers), value=diff)
        diff_cell.number_format = "+0.0;-0.0;0.0"
        if diff > 1:
            diff_cell.font = Font(color="FF0000", bold=True)
        elif diff < -1:
            diff_cell.font = Font(color="006600", bold=True)

        if i % 2 == 0:
            for col in range(1, len(headers) + 1):
                if ws.cell(row=r, column=col).fill == PatternFill():
                    ws.cell(row=r, column=col).fill = LIGHT_FILL

    _auto_width(ws)


def _build_overlap_analysis(wb: Workbook, analysis: dict):
    ws = wb.create_sheet("Overlap Analysis")

    overlaps = analysis.get("overlaps", [])

    ws.cell(row=1, column=1, value="Risk Overlap Analysis").font = Font(
        bold=True, size=14, color=DB_BLUE
    )

    row = 3
    headers = ["Risk Theme", "Tickers", "Combined Weight", "Avg Score", "Severity"]
    _header_row(ws, row, headers)

    severity_fills = {
        "CRITICAL": PatternFill("solid", fgColor="FF6B6B"),
        "HIGH": PatternFill("solid", fgColor="FFA500"),
        "MEDIUM": PatternFill("solid", fgColor="FFD700"),
        "LOW": PatternFill("solid", fgColor="90EE90"),
    }

    if not overlaps:
        ws.cell(row=row + 1, column=1, value="No significant risk overlaps detected.").font = Font(
            italic=True, color="888888"
        )
    else:
        for i, overlap in enumerate(overlaps):
            r = row + 1 + i
            ws.cell(row=r, column=1, value=overlap["factor_label"])
            ws.cell(row=r, column=2, value=", ".join(overlap["tickers"]))
            ws.cell(row=r, column=3, value=overlap["combined_weight"]).number_format = PCT_FORMAT
            ws.cell(row=r, column=4, value=overlap["avg_score"]).number_format = SCORE_FORMAT

            sev = overlap["severity"]
            sev_cell = ws.cell(row=r, column=5, value=sev)
            sev_cell.font = Font(bold=True)
            if sev in severity_fills:
                sev_cell.fill = severity_fills[sev]

    _auto_width(ws)


def _build_damodaran_reference(wb: Workbook, analysis: dict):
    ws = wb.create_sheet("Damodaran Reference")

    damodaran = analysis.get("damodaran_data", {})
    industry_map = analysis.get("industry_map", {})

    ws.cell(row=1, column=1, value="Damodaran Industry Reference Data").font = Font(
        bold=True, size=14, color=DB_BLUE
    )

    # Attribution
    ws.cell(
        row=2,
        column=1,
        value="Source: pages.stern.nyu.edu/~adamodar | Data as of January 2026 | Check source for updates",
    ).font = Font(italic=True, color="888888", size=9)

    # Industry mapping for this portfolio
    row = 4
    ws.cell(row=row, column=1, value="Portfolio Industry Mappings").font = BLUE_BOLD
    row += 1
    _header_row(ws, row, ["Ticker", "Damodaran Industry"])
    for i, (ticker, industry) in enumerate(industry_map.items()):
        r = row + 1 + i
        ws.cell(row=r, column=1, value=ticker).font = BLUE_FONT
        ws.cell(row=r, column=2, value=industry)

    # Full Damodaran table
    row = row + len(industry_map) + 3
    ws.cell(row=row, column=1, value="Full Damodaran Industry Data").font = BLUE_BOLD
    row += 1
    headers = ["Industry", "# Firms", "Levered Beta", "Unlevered Beta", "D/E Ratio", "Eff. Tax Rate"]
    _header_row(ws, row, headers)

    for i, (industry, data) in enumerate(sorted(damodaran.items())):
        if not isinstance(data, dict):
            continue
        r = row + 1 + i
        ws.cell(row=r, column=1, value=industry)
        ws.cell(row=r, column=2, value=data.get("n_firms", ""))
        ws.cell(row=r, column=3, value=data.get("levered_beta", "")).number_format = "0.00"
        ws.cell(row=r, column=4, value=data.get("unlevered_beta", "")).number_format = "0.00"
        ws.cell(row=r, column=5, value=data.get("de_ratio", "")).number_format = "0.00%"
        ws.cell(row=r, column=6, value=data.get("eff_tax_rate", "")).number_format = "0.00%"

        if i % 2 == 0:
            for col in range(1, len(headers) + 1):
                ws.cell(row=r, column=col).fill = LIGHT_FILL

    _auto_width(ws)
